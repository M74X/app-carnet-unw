"""Reporte de códigos repetidos e inválidos de la hoja BASE.

Uso: python scripts/repetidos.py data/archivo/<original>.xlsx
"""
import sys
import unicodedata
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

if len(sys.argv) != 2:
    sys.exit("Uso: python scripts/repetidos.py data/archivo/<original>.xlsx")

RAIZ    = Path(__file__).resolve().parent.parent
ARCHIVO = Path(sys.argv[1])
HOJA    = "BASE CARNÉ 2026-1"
SALIDA  = RAIZ / "reportes" / f"codigos_repetidos_{date.today().isoformat()}.xlsx"
SALIDA.parent.mkdir(exist_ok=True)


def norm(s):
    return unicodedata.normalize("NFD", str(s).strip().upper()).encode("ascii", "ignore").decode()


# --- Cargar ---
df = pd.read_excel(ARCHIVO, sheet_name=HOJA, dtype=str)
df.columns = [c.strip() for c in df.columns]
df.insert(0, "Fila Excel", df.index + 2)  # +1 encabezado, +1 base 1
df = df[df["Codigo"].notna()].copy()
df["Codigo"] = df["Codigo"].str.strip()
df["_estado"] = df["ESTADO"].map(lambda x: norm(x) if pd.notna(x) else "")

columnas = ["Fila Excel", "Codigo", "Primer Apellido", "Segundo Apellido", "Nombres",
            "Carrera", "Bloque", "fecha", "ESTADO", "Observaciones"]
columnas = [c for c in columnas if c in df.columns]

# --- Repetidos ---
rep = df[df.duplicated("Codigo", keep=False)].copy()
grupos = rep.groupby("Codigo")
primera = grupos["_estado"].transform("first")
rep["Estados distintos"] = grupos["_estado"].transform("nunique").gt(1).map({True: "SÍ", False: "NO"})
rep["La app muestra hoy"] = primera.map(lambda e: "ENTREGADO" if e == "ENTREGADO" else "PENDIENTE")
rep["Error en la app"] = (
    (primera != "ENTREGADO") & grupos["_estado"].transform(lambda s: (s == "ENTREGADO").any())
).map({True: "SÍ", False: "NO"})
rep = rep.sort_values(["Error en la app", "Estados distintos", "Codigo", "Fila Excel"],
                      ascending=[False, False, True, True])
rep = rep[columnas + ["Estados distintos", "La app muestra hoy", "Error en la app"]]

# --- Códigos inválidos (no tienen 10 caracteres) ---
inv = df[df["Codigo"].str.len() != 10][columnas].copy()

# --- Resumen ---
resumen = pd.DataFrame([
    ("Códigos repetidos",                                 rep["Codigo"].nunique()),
    ("Filas involucradas",                                len(rep)),
    ("Códigos con estados distintos entre filas",         rep.loc[rep["Estados distintos"] == "SÍ", "Codigo"].nunique()),
    ("Entregados que la app muestra como pendiente",      rep.loc[rep["Error en la app"] == "SÍ", "Codigo"].nunique()),
    ("Códigos inválidos (longitud distinta de 10)",       len(inv)),
], columns=["Indicador", "Cantidad"])

# --- Guardar ---
rojo    = PatternFill("solid", fgColor="F8D7DA")
naranja = PatternFill("solid", fgColor="FFF3CD")

with pd.ExcelWriter(SALIDA, engine="openpyxl") as xw:
    resumen.to_excel(xw, sheet_name="Resumen", index=False)
    rep.to_excel(xw, sheet_name="Repetidos", index=False)
    inv.to_excel(xw, sheet_name="Códigos inválidos", index=False)

    for ws in xw.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            ancho = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 45)

    ws = xw.book["Repetidos"]
    enc = [c.value for c in ws[1]]
    i_err, i_dist = enc.index("Error en la app"), enc.index("Estados distintos")
    for row in ws.iter_rows(min_row=2):
        fill = rojo if row[i_err].value == "SÍ" else naranja if row[i_dist].value == "SÍ" else None
        if fill:
            for c in row:
                c.fill = fill

print(resumen.to_string(index=False))
print(f"Listo — {SALIDA.relative_to(RAIZ)} generado")
